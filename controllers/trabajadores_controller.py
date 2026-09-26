from helpers import clean_amount
import io
from fpdf import FPDF
from flask import send_file, Blueprint, render_template, session, redirect, url_for, flash, request
from models import db, Usuarios, Tarjeta, ProgramacionPagoTarjeta
from datetime import datetime
from decorators import login_required, admin_oficina_required
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
trabajadores_bp = Blueprint("trabajadores", __name__)

@trabajadores_bp.route("/oficina/trabajadores")
@login_required
@admin_oficina_required
def index():
    usuario = Usuarios.query.get(session["user_id"])
    return render_template("trabajadores/index.html", usuario=usuario)

@trabajadores_bp.route("/oficina/trabajadores/nomina")
@login_required
@admin_oficina_required
def nomina():
    usuario = Usuarios.query.get(session["user_id"])
    empleados = Usuarios.query.filter_by(rol="EMPLEADO").all()
    return render_template("trabajadores/nomina.html", usuario=usuario, empleados=empleados)

@trabajadores_bp.route("/oficina/trabajadores/tarjetas")
@login_required
@admin_oficina_required
def tarjetas():
    usuario = Usuarios.query.get(session["user_id"])
    tarjetas_lista = Tarjeta.query.all()
    pagos_programados = ProgramacionPagoTarjeta.query.filter(ProgramacionPagoTarjeta.estado != 'REALIZADO').order_by(ProgramacionPagoTarjeta.fecha_programada.asc()).all()
    total_programado = sum(float(p.monto or 0) for p in pagos_programados)
    
    return render_template("trabajadores/tarjetas.html", usuario=usuario, tarjetas=tarjetas_lista, pagos_programados=pagos_programados, total_programado=total_programado)

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/historial")
@login_required
@admin_oficina_required
def historial_tarjetas():
    usuario = Usuarios.query.get(session["user_id"])
    fecha_inicio = request.args.get("fecha_inicio")
    fecha_fin = request.args.get("fecha_fin")
    
    query = ProgramacionPagoTarjeta.query.filter_by(estado='REALIZADO')
    if fecha_inicio:
        try:
            fecha_ini_date = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
            query = query.filter(ProgramacionPagoTarjeta.fecha_programada >= fecha_ini_date)
        except ValueError:
            pass
    if fecha_fin:
        try:
            fecha_fin_date = datetime.strptime(fecha_fin, "%Y-%m-%d").date()
            query = query.filter(ProgramacionPagoTarjeta.fecha_programada <= fecha_fin_date)
        except ValueError:
            pass
            
    historial_pagos = query.order_by(ProgramacionPagoTarjeta.fecha_programada.desc()).all()
    
    return render_template("trabajadores/historial_tarjetas.html", usuario=usuario, historial_pagos=historial_pagos, filtro_inicio=fecha_inicio, filtro_fin=fecha_fin)

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/crear", methods=["POST"])
@login_required
@admin_oficina_required
def crear_tarjeta():
    try:
        nombre = request.form.get("nombre")
        numero = request.form.get("numero")
        
        if nombre and numero:
            nueva_tarjeta = Tarjeta(nombre=nombre, numero=numero)
            db.session.add(nueva_tarjeta)
            db.session.commit()
            flash("Tarjeta creada correctamente", "success")
        else:
            flash("Nombre y número son obligatorios", "warning")
            
    except Exception as e:
        db.session.rollback()
        flash(f"Error al crear la tarjeta: {str(e)}", "danger")
        
    return redirect(url_for("trabajadores.tarjetas"))

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/editar", methods=["POST"])
@login_required
@admin_oficina_required
def editar_tarjeta():
    try:
        id_tarjeta = request.form.get("id_tarjeta")
        nombre = request.form.get("nombre")
        numero = request.form.get("numero")
        
        tarjeta = Tarjeta.query.get_or_404(id_tarjeta)
        
        if nombre and numero:
            tarjeta.nombre = nombre
            tarjeta.numero = numero
            db.session.commit()
            flash("Tarjeta actualizada con éxito", "success")
        else:
            flash("Nombre y número son obligatorios", "warning")
            
    except Exception as e:
        db.session.rollback()
        flash(f"Error al actualizar la tarjeta: {str(e)}", "danger")
        
    return redirect(url_for("trabajadores.tarjetas"))

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/eliminar/<int:id>", methods=["POST"])
@login_required
@admin_oficina_required
def eliminar_tarjeta(id):
    try:
        tarjeta = Tarjeta.query.get_or_404(id)
        ProgramacionPagoTarjeta.query.filter_by(tarjeta_id=id).delete()
        db.session.delete(tarjeta)
        db.session.commit()
        flash("Tarjeta eliminada correctamente.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Error al eliminar la tarjeta: {str(e)}", "danger")
        
    return redirect(url_for("trabajadores.tarjetas"))

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/programar-pago", methods=["POST"])
@login_required
@admin_oficina_required
def programar_pago_tarjeta():
    try:
        tarjeta_id = request.form.get("tarjeta_id")
        monto = clean_amount(request.form.get("monto"))
        fecha_str = request.form.get("fecha_programada")
        concepto = request.form.get("concepto")
        cuenta_origen = request.form.get("cuenta_origen")
        
        if not (tarjeta_id and monto and fecha_str):
            flash("Todos los campos obligatorios deben ser completados.", "warning")
            return redirect(url_for("trabajadores.tarjetas"))
            
        fecha_prog = datetime.strptime(fecha_str, "%Y-%m-%d").date()
        
        nuevo_pago = ProgramacionPagoTarjeta(
            tarjeta_id=int(tarjeta_id),
            monto=float(monto),
            fecha_programada=fecha_prog,
            concepto=concepto,
            cuenta_origen=cuenta_origen
        )
        
        db.session.add(nuevo_pago)
        db.session.commit()
        flash("Pago programado correctamente", "success")
        
    except Exception as e:
        db.session.rollback()
        flash(f"Error al programar el pago: {str(e)}", "danger")
        
    return redirect(url_for("trabajadores.tarjetas"))

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/editar-pago", methods=["POST"])
@login_required
@admin_oficina_required
def editar_pago_programado():
    try:
        pago_id = request.form.get("pago_id")
        fecha_str = request.form.get("fecha_programada")
        monto = clean_amount(request.form.get("monto"))
        cuenta_origen = request.form.get("cuenta_origen")
        concepto = request.form.get("concepto")
        
        if not (pago_id and fecha_str and monto):
            flash("Faltan campos obligatorios.", "warning")
            return redirect(url_for("trabajadores.tarjetas"))
            
        pago = ProgramacionPagoTarjeta.query.get_or_404(pago_id)
        
        fecha_prog = datetime.strptime(fecha_str, "%Y-%m-%d").date()
        
        pago.fecha_programada = fecha_prog
        pago.monto = float(monto)
        pago.cuenta_origen = cuenta_origen
        pago.concepto = concepto
        
        db.session.commit()
        flash("Pago programado actualizado correctamente.", "success")
        
    except Exception as e:
        db.session.rollback()
        flash(f"Error al actualizar el pago programado: {str(e)}", "danger")
        
    return redirect(url_for("trabajadores.tarjetas"))
@trabajadores_bp.route("/oficina/trabajadores/tarjetas/eliminar-pago/<int:id>", methods=["POST"])
@login_required
@admin_oficina_required
def eliminar_pago_tarjeta(id):
    try:
        pago = ProgramacionPagoTarjeta.query.get_or_404(id)
        db.session.delete(pago)
        db.session.commit()
        flash("Pago programado eliminado correctamente", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Error al eliminar el pago: {str(e)}", "danger")
        
    return redirect(url_for("trabajadores.tarjetas"))

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/marcar-pagado/<int:id>", methods=["POST"])
@login_required
@admin_oficina_required
def marcar_pagado_tarjeta(id):
    try:
        pago = ProgramacionPagoTarjeta.query.get_or_404(id)
        forma_pago = request.form.get("forma_pago")
        pago.estado = 'REALIZADO'
        if forma_pago:
            pago.forma_pago = forma_pago
        db.session.commit()
        flash("Pago marcado como REALIZADO correctamente", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Error al marcar el pago: {str(e)}", "danger")
    return redirect(url_for("trabajadores.tarjetas"))

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/exportar-pdf")
@login_required
@admin_oficina_required
def exportar_pdf_tarjetas():
    try:
        pagos = ProgramacionPagoTarjeta.query.order_by(ProgramacionPagoTarjeta.fecha_programada.asc()).all()
        
        pdf = FPDF(orientation='L', format="A4")
        pdf.add_page()
        pdf.set_margins(15, 15, 15)
        pdf.set_font("Helvetica", style="B", size=16)
        pdf.cell(0, 10, "REPORTE DE PROGRAMACION DE PAGOS - TARJETAS", align="C")
        pdf.ln(15)
        
        pdf.set_font("Helvetica", style="B", size=10)
        pdf.set_fill_color(220, 220, 220)
        pdf.cell(30, 10, "FECHA", border=1, fill=True, align="C")
        pdf.cell(50, 10, "TARJETA", border=1, fill=True, align="C")
        pdf.cell(35, 10, "MONTO", border=1, fill=True, align="C")
        pdf.cell(50, 10, "CUENTA ORIGEN", border=1, fill=True, align="C")
        pdf.cell(70, 10, "CONCEPTO", border=1, fill=True, align="C")
        pdf.cell(30, 10, "ESTADO", border=1, fill=True, align="C")
        pdf.ln(10)
        
        pdf.set_font("Helvetica", size=9)
        total_monto = 0.0
        
        for p in pagos:
            tarjeta_obj = getattr(p, 'tarjeta', None)
            tarjeta_nombre = tarjeta_obj.nombre if (tarjeta_obj and hasattr(tarjeta_obj, 'nombre') and tarjeta_obj.nombre) else 'Desconocida'
            
            t_n = (str(tarjeta_nombre)[:20] + '..') if len(str(tarjeta_nombre)) > 20 else str(tarjeta_nombre)
            
            cuenta_origen = getattr(p, 'cuenta_origen', None) or '-'
            c_o = (str(cuenta_origen)[:20] + '..') if len(str(cuenta_origen)) > 20 else str(cuenta_origen)
            
            concepto = getattr(p, 'concepto', None) or '-'
            con = (str(concepto)[:35] + '..') if len(str(concepto)) > 35 else str(concepto)
            
            f_prog = getattr(p, 'fecha_programada', None)
            if f_prog and hasattr(f_prog, 'strftime'):
                fecha_str = f_prog.strftime('%d/%m/%Y')
            else:
                fecha_str = str(f_prog or '-')
                
            try:
                monto_num = float(getattr(p, 'monto', 0) or 0)
            except (ValueError, TypeError):
                monto_num = 0.0
            total_monto += monto_num
            
            estado_str = str(getattr(p, 'estado', '') or 'Programado')
            
            pdf.cell(30, 8, fecha_str, border=1, align="C")
            pdf.cell(50, 8, t_n, border=1)
            pdf.cell(35, 8, f"$ {monto_num:,.2f}", border=1, align="R")
            pdf.cell(50, 8, c_o, border=1)
            pdf.cell(70, 8, con, border=1)
            pdf.cell(30, 8, estado_str, border=1, align="C")
            pdf.ln(8)
            
        pdf.set_font("Helvetica", style="B", size=9)
        pdf.set_fill_color(240, 240, 240)
        pdf.cell(80, 8, "SUMA TOTAL", border=1, fill=True, align="R")
        pdf.cell(35, 8, f"$ {total_monto:,.2f}", border=1, fill=True, align="R")
        pdf.cell(150, 8, "", border=1, fill=True)
        pdf.ln(12)
        
        pdf.set_font("Helvetica", style="I", size=8)
        pdf.cell(0, 5, "Documento generado automaticamente.", align="L")
        
        byte_output = pdf.output()
        if isinstance(byte_output, str):
            byte_output = byte_output.encode('latin1')
        else:
            byte_output = bytes(byte_output)
            
        return send_file(
            io.BytesIO(byte_output),
            download_name="Reporte_Pagos_Tarjetas.pdf",
            mimetype="application/pdf"
        )
    except Exception as e:
        import traceback
        print(f"Error al exportar PDF de tarjetas: {e}")
        print(traceback.format_exc())
        flash(f"Error al generar el archivo PDF: {str(e)}", "danger")
        return redirect(url_for("trabajadores.tarjetas"))

@trabajadores_bp.route("/oficina/trabajadores/tarjetas/exportar-pdf-historial")
@login_required
@admin_oficina_required
def exportar_pdf_historial():
    fecha_inicio = request.args.get("fecha_inicio")
    fecha_fin = request.args.get("fecha_fin")
    
    query = ProgramacionPagoTarjeta.query.filter_by(estado='REALIZADO')
    if fecha_inicio:
        try:
            fecha_ini_date = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
            query = query.filter(ProgramacionPagoTarjeta.fecha_programada >= fecha_ini_date)
        except ValueError:
            pass
    if fecha_fin:
        try:
            fecha_fin_date = datetime.strptime(fecha_fin, "%Y-%m-%d").date()
            query = query.filter(ProgramacionPagoTarjeta.fecha_programada <= fecha_fin_date)
        except ValueError:
            pass
            
    historial = query.order_by(ProgramacionPagoTarjeta.fecha_programada.desc()).all()
    
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()
    elements.append(Paragraph("Historial de Pagos Realizados (Tarjetas)", styles['Title']))
    elements.append(Spacer(1, 20))
    
    data = [["FECHA", "TARJETA", "MONTO", "CUENTA ORIGEN", "CONCEPTO", "ESTADO"]]
    for p in historial:
        tarjeta_nombre = p.tarjeta.nombre if p.tarjeta and p.tarjeta.nombre else 'Desconocida'
        data.append([
            p.fecha_programada.strftime('%d/%m/%Y'),
            tarjeta_nombre[:20] + '..' if len(tarjeta_nombre) > 20 else tarjeta_nombre,
            f"$ {p.monto:,.2f}",
            (p.cuenta_origen[:20] + '..') if p.cuenta_origen and len(p.cuenta_origen) > 20 else (p.cuenta_origen or '-'),
            (p.concepto[:35] + '..') if p.concepto and len(p.concepto) > 35 else (p.concepto or '-'),
            str(p.estado)
        ])
        
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.whitesmoke),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    elements.append(t)
    doc.build(elements)
    
    output.seek(0)
    return send_file(
        output,
        download_name="Historial_Pagos_Tarjetas.pdf",
        mimetype="application/pdf",
        as_attachment=True
    )
